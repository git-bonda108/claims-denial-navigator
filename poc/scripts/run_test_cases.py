"""Run the business test cases against a Denial Navigator app and write an Excel test log.

  python3 scripts/run_test_cases.py https://<your-app>.azurecontainerapps.io [--out results.xlsx]

Every case has plain-language steps a business user can follow in the web app, the expected
result, and an automated run of the same steps through the app's API. The workbook records the
captured result, PASS/FAIL and response time, plus empty columns for a person to validate by hand.
The run starts and ends with "Reset demo data", so the app is clean for a live demo afterwards.
"""
from __future__ import annotations

import argparse
import json
import os
import time
import urllib.error
import urllib.request
import uuid
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "DM Sans"  # Alif deck fonts: Anthropic Sans on the presenter Mac, DM Sans for everyone else
NAVY, BLUE, TEAL, GREEN, CORAL, AMBER, LIGHT = "0B1F3A", "1F4E8C", "00A3C4", "2E9E6A", "E4572E", "F2A541", "F3F4F6"


class App:
    def __init__(self, base: str):
        self.base = base.rstrip("/")

    def call(self, method: str, path: str, body=None, timeout=180, raw: bytes | None = None, ctype: str | None = None):
        data = json.dumps(body).encode() if body is not None else raw
        headers = {"Content-Type": "application/json"} if body is not None else ({"Content-Type": ctype} if ctype else {})
        req = urllib.request.Request(self.base + path, data=data, method=method, headers=headers)
        t0 = time.time()
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                out = json.loads(r.read().decode())
        except urllib.error.HTTPError as e:
            out = {"error": f"HTTP {e.code}", "detail": e.read().decode()[:200]}
        return out, time.time() - t0

    def fetch(self, path: str, timeout=120):
        """GET a page or file: (status, headers, body bytes, seconds)."""
        t0 = time.time()
        try:
            with urllib.request.urlopen(self.base + path, timeout=timeout) as r:
                return r.status, dict(r.headers), r.read(), time.time() - t0
        except urllib.error.HTTPError as e:
            return e.code, dict(e.headers), e.read(), time.time() - t0

    def wait(self, limit=300):
        """Wait until the app is really serving: right after a deploy or restart Azure answers 503 for a while,
        and the first real request also generates the demo dataset."""
        t0, said = time.time(), False
        while time.time() - t0 < limit:
            try:
                meta = self.call("GET", "/api/meta", timeout=120)[0]
                if isinstance(meta, dict) and "agent_mode" in meta:
                    return meta
            except (urllib.error.URLError, TimeoutError, ValueError):
                pass
            if not said:
                print("Waiting for the app to start (up to 5 minutes after a deploy or restart)...")
                said = True
            time.sleep(5)
        raise SystemExit("The app did not become ready within 5 minutes - check it in the browser, then run again.")

    def ingest(self, prefix: str):
        resp = self.call("GET", "/api/demo-files")[0]
        name = next((f for f in resp.get("files", []) if f.startswith(prefix)), None)
        if not name:
            raise RuntimeError(f"No live-demo file starting with {prefix} (app returned: {str(resp)[:120]})")
        r, dt = self.call("POST", f"/api/demo-files/{name}/ingest")
        r["file"] = name
        return r, dt

    def decisions(self, claim_id: str, explain=False):
        det, dt = self.call("GET", f"/api/claims/{claim_id}{'?explain=true' if explain else ''}")
        return [d for d in det.get("decisions", []) if d.get("workable")], det, dt


ACTION = {"APPEAL_LEVEL_1": "Appeal (level 1)", "CORRECT_AND_RESUBMIT": "Correct & resubmit", "BILL_OTHER_PAYER": "Bill other payer",
          "SUBMIT_NEW_CLAIM": "Submit new claim", "WRITE_OFF": "Write off", "BILL_PATIENT": "Bill patient"}
act = lambda a: ACTION.get(a, a)


# ---------------------------------------------------------------- test cases
def tc_version(app, ctx):
    status, headers, body, dt = app.fetch("/")
    meta = app.call("GET", "/api/meta")[0]
    build = meta.get("build", "")
    html = body.decode("utf-8", "replace")
    cache = headers.get("Cache-Control", headers.get("cache-control", ""))
    ok = status == 200 and bool(build) and f"app.js?v={build}" in html and "no-store" in cache
    ctx["build"] = build
    return ok, (f"Page HTTP {status}. Version {build or 'missing'}; page loads app.js?v={build}: {f'app.js?v={build}' in html}. "
                f"Cache-Control: {cache or 'none'} (browsers always fetch the latest version). AI mode: {meta.get('agent_mode')}."), dt


def tc_reset(app, ctx):
    r, dt = app.call("POST", "/api/demo/reset", timeout=300)
    s = app.call("GET", "/api/summary")[0]
    live = [f["name"] for f in s.get("files", []) if "LIVE" in f["name"]]
    ok = r.get("claims", 0) > 1000 and s.get("completed") == 0 and not s.get("alerts") and not live
    return ok, (f"Reloaded {r.get('claims')} claims and {r.get('remits')} remittances. Resolved: {s.get('completed')}. "
                f"Teams alerts: {len(s.get('alerts', []))}. Live files in Recent files: {len(live)}."), dt


def tc_dashboard(app, ctx):
    s, dt = app.call("GET", "/api/summary")
    ok = (0.15 <= s.get("denial_rate", 0) <= 0.21 and s.get("open_items", 0) > 100 and s.get("noise_filtered", 0) > 1000
          and s.get("denied_amount", 0) > 0 and s.get("completed") == 0 and s.get("by_category") and s.get("by_payer"))
    return ok, (f"{s.get('claims'):,} claims, denial rate {s.get('denial_rate', 0):.1%}. Open denied ${s.get('denied_amount', 0):,.0f} "
                f"across {s.get('open_items')} items; at risk ≤30 days ${s.get('at_risk_30d', 0):,.0f}. "
                f"{s.get('noise_filtered', 0):,} routine adjustments filtered. Resolved {s.get('completed')}. "
                f"Top root cause: {max(s.get('by_category', {'-': 0}), key=s.get('by_category', {'-': 0}).get)}."), dt


def tc_worklist(app, ctx):
    items, dt = app.call("GET", "/api/worklist")
    items = items.get("items", [])
    coding, _ = app.call("GET", "/api/worklist?queue=CODING")
    crit, _ = app.call("GET", "/api/worklist?band=CRITICAL")
    scores = [i["score"] for i in items]
    ok = (items and scores == sorted(scores, reverse=True) and coding["items"] and {i["queue"] for i in coding["items"]} == {"CODING"}
          and all(i["band"] == "CRITICAL" for i in crit["items"]) and all(i.get("action") and i.get("queue") for i in items))
    return ok, (f"{len(items)} open items, sorted by score ({scores[0]} → {scores[-1]}). Queue filter Coding: {len(coding['items'])} items, "
                f"all Coding: {ok}. Urgency filter Critical: {len(crit['items'])} items. Every item has a decision and a team queue."), dt


def tc_smm(app, ctx):
    r, dt = app.ingest("835_LIVE_SMM")
    ctx["smm"] = r
    s = app.call("GET", "/api/summary")[0]
    alerts = [a["card"]["body"][0]["text"] for a in s.get("alerts", [])]
    top = (s.get("files") or [{}])[0]
    ok = (r.get("status") == "PROCESSED" and r.get("denials") == 2 and r.get("alerts", 0) == 2 and len(alerts) == 2
          and top.get("name") == r.get("file"))
    return ok, (f"File {r.get('status')} in {dt:.1f}s: {r.get('claims')} claims, {r.get('denials')} denials decided, "
                f"{r.get('alerts')} Teams alerts ({'; '.join(alerts)}). Top of Recent files: {top.get('name')}. New items: "
                + ", ".join(f"{i['claim_id']} → {act(i['action'])} ${i['amount']:,.0f}" for i in r.get("new_items", []))), dt


def tc_prior_auth(app, ctx):
    cid = next(i["claim_id"] for i in ctx["smm"]["new_items"] if i["action"] == "APPEAL_LEVEL_1")
    ds, det, dt = app.decisions(cid)
    d = ds[0]
    ex = d.get("explanation") or {}
    letter, dt2 = app.call("POST", f"/api/decisions/{d['id']}/letter")
    auth = (det.get("ehr", {}).get("auth_on_file") or {}).get("number", "")
    appr, _ = app.call("POST", f"/api/decisions/{d['id']}/feedback", {"verdict": "APPROVE", "note": "Test case"})
    s = app.call("GET", "/api/summary")[0]
    ok = (d["action"] == "APPEAL_LEVEL_1" and d["queue"] == "UTILIZATION_REVIEW" and auth and auth in d["root_cause"]
          and ex.get("citations") and auth in letter.get("text", "") and cid in letter.get("text", "")
          and appr.get("status") == "DONE" and s.get("completed") == 1)
    return ok, (f"{cid} CO-{d['carc']} ${d['amount']:,.0f}: {act(d['action'])} → Utilization Review. Cause: {d['root_cause']}. "
                f"Due {d['deadline']['due_date']} ({d['deadline']['days_left']} days). Cited: {', '.join(c['id'] for c in ex.get('citations', []))}. "
                f"Explanation by {ex.get('by')}. Letter by {letter.get('by')} in {dt2:.1f}s contains claim and auth number: "
                f"{auth in letter.get('text', '') and cid in letter.get('text', '')}. Approved → {appr.get('status')}; Resolved now {s.get('completed')}."), dt + dt2


def tc_deadline(app, ctx):
    cid = next(i["claim_id"] for i in ctx["smm"]["new_items"] if i["action"] == "CORRECT_AND_RESUBMIT")
    ds, det, dt = app.decisions(cid)
    d = ds[0]
    reply = app.call("POST", "/api/chat", {"message": "What is due this week?"})[0]
    ok = d["action"] == "CORRECT_AND_RESUBMIT" and d["queue"] == "CODING" and d["deadline"]["days_left"] <= 7 and d["deadline"]["band"] == "CRITICAL"
    return ok, (f"{cid} CO-{d['carc']} ${d['amount']:,.0f}: {act(d['action'])} → Coding. Cause: {d['root_cause']}. "
                f"Due {d['deadline']['due_date']}: {d['deadline']['days_left']} days, {d['deadline']['band'].title()}. "
                f"Assistant 'due this week' mentions it: {cid in reply.get('reply', '')}."), dt


def tc_medicare(app, ctx):
    r, dt = app.ingest("835_LIVE_MCR")
    ctx["mcr"] = r
    cid = next(i["claim_id"] for i in r["new_items"] if i["action"] == "CORRECT_AND_RESUBMIT")
    ds, det, _ = app.decisions(cid)
    d = next(x for x in ds if x["carc"] == "50")
    fail = next((c for c in d["checks"] if c["result"] == "FAIL"), {})
    found = next((c for c in d["checks"] if c["result"] == "PASS"), {})
    by = {x.get("explanation", {}).get("by", "") for x in ds}
    ok = (r.get("status") == "PROCESSED" and r.get("denials") == 6 and d["action"] == "CORRECT_AND_RESUBMIT" and d["queue"] == "CODING"
          and fail and found and d["deadline"]["type"] == "REOPENING")
    return ok, (f"File {r.get('status')} in {dt:.1f}s: {r.get('claims')} claims, {r.get('denials')} denials. {cid} CO-50 {' '.join(d['rarcs'])} "
                f"on {d['code']} ${d['amount']:,.0f}: {act(d['action'])} → Coding. Check failed: {fail.get('name')} ({fail.get('detail')}). "
                f"Found: {found.get('name')} ({found.get('detail')}). Deadline: {d['deadline']['rule']} ({d['deadline']['due_date']}). "
                f"Explanation by {', '.join(sorted(by))}."), dt


def tc_missing_npi(app, ctx):
    cid = next(i["claim_id"] for i in ctx["mcr"]["new_items"] if i["action"] == "SUBMIT_NEW_CLAIM")
    ds, det, dt = app.decisions(cid)
    d = ds[0]
    fail = next((c for c in d["checks"] if c["result"] == "FAIL"), {})
    ok = d["action"] == "SUBMIT_NEW_CLAIM" and d["queue"] == "BILLING" and "MA130" in d["rarcs"]
    return ok, (f"{cid} CO-16 {' '.join(d['rarcs'])}: {act(d['action'])} → Billing. Cause: {d['root_cause'][:110]}. "
                f"Check: {fail.get('name')}. Steps: {' / '.join(d['steps'][:2])}. No appeal recommended."), dt


def tc_other_denials(app, ctx):
    items = ctx["mcr"]["new_items"]
    cob = next(i["claim_id"] for i in items if i["action"] == "BILL_OTHER_PAYER")
    bun = [i["claim_id"] for i in items if i["action"] == "CORRECT_AND_RESUBMIT"][-1]
    a, _, dt = app.decisions(cob)
    b, _, dt2 = app.decisions(bun)
    b0 = next(x for x in b if x["carc"] == "97")
    ok = (all(x["action"] == "BILL_OTHER_PAYER" and x["queue"] == "REGISTRATION" for x in a)
          and b0["action"] == "CORRECT_AND_RESUBMIT" and b0["queue"] == "CODING")
    return ok, (f"{cob} CO-22 ({len(a)} lines, ${sum(x['amount'] for x in a):,.0f}): {act(a[0]['action'])} → Registration. "
                f"Cause: {a[0]['root_cause']}. {bun} CO-97 {' '.join(b0['rarcs'])} on {b0['code']}: {act(b0['action'])} → Coding. "
                f"Cause: {b0['root_cause']}."), dt + dt2


def tc_explain(app, ctx):
    wl = app.call("GET", "/api/worklist")[0]["items"]
    live = {i["claim_id"] for i in ctx["smm"]["new_items"] + ctx["mcr"]["new_items"]}
    cid = next(i["claim_id"] for i in wl if i["claim_id"] not in live)
    before, _, _ = app.decisions(cid)
    after, _, dt = app.decisions(cid, explain=True)
    ex = after[0].get("explanation") or {}
    mode = app.call("GET", "/api/meta")[0].get("agent_mode")
    ok = bool(ex.get("narrative")) and after[0]["action"] == before[0]["action"] and (mode != "azure" or ex.get("by", "").startswith("azure"))
    return ok, (f"{cid}: explanation by {ex.get('by')} in {dt:.1f}s. Rule decision unchanged ({act(after[0]['action'])}). "
                f"Starts: {ex.get('narrative', '')[:150]!r}"), dt


def tc_override(app, ctx):
    wl = app.call("GET", "/api/worklist")[0]["items"]
    target = next(i for i in wl if i["action"] != "WRITE_OFF")
    k0 = len(app.call("GET", "/api/knowledge")[0]["cards"])
    bad, _ = app.call("POST", f"/api/decisions/{target['id']}/feedback", {"verdict": "OVERRIDE", "action": "NOT_AN_ACTION", "note": "x"})
    o, dt = app.call("POST", f"/api/decisions/{target['id']}/feedback",
                     {"verdict": "OVERRIDE", "action": "WRITE_OFF", "note": "Test case: payer confirmed by phone"})
    k1 = len(app.call("GET", "/api/knowledge")[0]["cards"])
    s = app.call("GET", "/api/summary")[0]
    ok = "error" in bad and o.get("status") == "DONE" and k1 == k0 + 1 and s.get("completed") == 2
    return ok, (f"Invalid action rejected: {'error' in bad}. Override of {target['claim_id']} ({act(target['action'])} → Write off) saved: "
                f"{o.get('status')}. Policy library {k0} → {k1} cards (new precedent). Resolved now {s.get('completed')}."), dt


def tc_assistant(app, ctx):
    cid = next(i["claim_id"] for i in ctx["smm"]["new_items"] if i["action"] == "APPEAL_LEVEL_1")
    qs = ["What is due this week?", f"Why was {cid} denied?", "What is the Medicare timely filing limit?", "Write me a poem about cats"]
    out, total = [], 0.0
    for q in qs:
        r, dt = app.call("POST", "/api/chat", {"message": q})
        out.append(r)
        total += dt
    ok = (all(r.get("reply") for r in out) and cid in out[1]["reply"] and ("12" in out[2]["reply"] or "one" in out[2]["reply"].lower())
          and not any(w in out[3]["reply"].lower() for w in ("meow", "whisker", "purr")))
    return ok, (f"All 4 answered by {out[0].get('by')} ({total:.1f}s total). Due this week: {out[0]['reply'][:90]!r}. "
                f"Why {cid}: {out[1]['reply'][:90]!r}. Timely filing: {out[2]['reply'][:70]!r}. Off-topic: {out[3]['reply'][:70]!r}."), total


def tc_files(app, ctx):
    name = app.call("GET", "/api/demo-files")[0]["files"][0]
    status, _, body, dt = app.fetch(f"/api/demo-files/{name}")
    boundary = uuid.uuid4().hex
    raw = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"broken.txt\"\r\nContent-Type: text/plain\r\n\r\n"
           "this is not an X12 file\r\n" f"--{boundary}--\r\n").encode()
    q, dt2 = app.call("POST", "/api/ingest", raw=raw, ctype=f"multipart/form-data; boundary={boundary}")
    ok = status == 200 and body.startswith(b"ISA") and q.get("status") == "QUARANTINED"
    return ok, (f"'view raw' on {name}: HTTP {status}, starts with {body[:3].decode()} (X12). Uploaded a broken file: {q.get('status')} "
                f"- {q.get('error')}. The app kept running."), dt + dt2


def tc_proof(app, ctx):
    e, dt = app.call("GET", "/api/evals")
    s = app.call("GET", "/api/summary")[0]
    acc = e.get("accuracy", {})
    ok = acc.get("denial", 0) >= 0.95 and acc.get("noise", 0) >= 0.99 and e.get("must_not_violations") == []
    return ok, (f"Denial decisions correct {acc.get('denial', 0):.1%} ({e.get('counts', {}).get('denial')} labelled). "
                f"Noise filtered correctly {acc.get('noise', 0):.1%}. Unsafe actions: {len(e.get('must_not_violations', []))} "
                f"(human overrides are not counted against the system). Deadlines computed: {e.get('deadlines_computed')}. "
                f"Ranking AUC {s.get('model_card', {}).get('test_auc')}."), dt


def tc_explainer(app, ctx):
    res, total = [], 0.0
    for path, kind in (("/explainer", b"<"), ("/explainer.gif", b"GIF8"), ("/explainer.mp4", None)):
        status, headers, body, dt = app.fetch(path)
        total += dt
        good = status == 200 and len(body) > 1000 and (body.startswith(kind) if kind else b"ftyp" in body[:16])
        res.append((path, status, len(body), good))
    ok = all(r[3] for r in res)
    return ok, " · ".join(f"{p}: HTTP {s}, {n / 1_000_000:.1f} MB" if n > 100_000 else f"{p}: HTTP {s}, {n // 1000} KB" for p, s, n, _ in res), total


def tc_clean(app, ctx):
    ok, actual, dt = tc_reset(app, ctx)
    return ok, actual + " Ready for a live demo.", dt


CASES = [
    ("TC-01", "Setup", "Browsers get the latest version", "Will everyone see the newest screens, not an old cached copy?",
     "1. Open the app link.\n2. Look at the bottom of the left sidebar.",
     "Sidebar shows 'Version' followed by the code in the Actual result column. The Remittance inbox has a 'Reset demo data' button.", tc_version),
    ("TC-02", "Setup", "Reset demo data", "Can we put the demo back to a clean start without IT?",
     "1. Remittance inbox → Reset demo data → OK.\n2. Open Dashboard.",
     "Takes about 10 seconds. Dashboard: Resolved 0, Teams alerts 'No alerts yet'. Live files show a plain 'Ingest' button.", tc_reset),
    ("TC-03", "Dashboard", "Denials at a glance", "Does the dashboard show the money and urgency at a glance?",
     "1. Open Dashboard.\n2. Read the 5 tiles, the urgency bar and the 4 charts.",
     "About 1,200 claims, denial rate 15-21%. Open denied dollars, at-risk ≤30 days, over 3,000 routine adjustments filtered, Resolved 0. "
     "Charts by root cause, payer and team queue.", tc_dashboard),
    ("TC-04", "Worklist", "Ranked and filtered worklist", "Does each team see only its own work, most valuable and urgent first?",
     "1. Open Worklist.\n2. Set Queue = Coding.\n3. Set Queue = All queues, Urgency = Critical.",
     "Rows sorted by score (highest first). Coding filter shows only Coding rows. Critical filter shows only critical deadlines.", tc_worklist),
    ("TC-05", "Ingestion", "A new remittance arrives (Summit Mutual)", "Does a payer file turn into decided denials and alerts, without anyone reading it?",
     "1. Remittance inbox → click Ingest on 835_LIVE_SMM_… once.\n2. Watch the button count 'Processing… 1s, 2s'.\n3. Open Dashboard → Teams alerts.",
     "Done within about 5 seconds: '2 denials decided, 2 Teams alerts'. File gets an 'Ingested ✓' badge and tops Recent files. "
     "Alerts: High-value denial (PRH-10222, $2,840) and Deadline alert (PRH-10223).", tc_smm),
    ("TC-06", "Decision", "Prior-auth denial: explain, cite, letter, approve", "Can a biller understand a $2,840 denial and act on it in one screen?",
     "1. Dashboard → Teams alerts → 'Open claim' on PRH-10222.\n2. Read the decision card.\n3. Draft letter → read → close.\n4. Approve plan.",
     "Appeal (level 1) → Utilization Review. Cause names authorization AUTH-… that was on file but not on the claim. Policy citation and "
     "60-day deadline shown. Letter has the claim and auth numbers. Approve → status DONE, Dashboard Resolved 1.", tc_prior_auth),
    ("TC-07", "Decision", "Deadline alert: coding fix due in 5 days", "Does it catch money about to be lost to a deadline?",
     "1. Open PRH-10223 (Teams alert or Worklist).\n2. Read Deadline and Checks.",
     "Correct & resubmit → Coding. Cause: billed diagnosis inconsistent with the knee injection side; the chart documents the right code. "
     "Deadline 5 days, Critical (red).", tc_deadline),
    ("TC-08", "Ingestion", "Medicare medical-necessity denial", "Does it find the real fix instead of writing the money off?",
     "1. Remittance inbox → Ingest 835_LIVE_MCR_… once.\n2. Click PRH-10224 in the result.\n3. Read Checks and Steps.",
     "Done within about 10 seconds: 4 claims, 6 denials. PRH-10224: Correct & resubmit → Coding. ✗ billed diagnosis not covered, "
     "✓ covered diagnosis found in the chart. Deadline: reopening within 1 year.", tc_medicare),
    ("TC-09", "Decision", "Unprocessable claim (missing provider ID)", "Does it avoid wasting an appeal on a claim with no appeal rights?",
     "1. Open PRH-10227.\n2. Read Checks and Steps.",
     "Submit new claim → Billing. Rendering provider NPI missing. Steps: add the NPI, send a new claim. No appeal.", tc_missing_npi),
    ("TC-10", "Decision", "Other payer first; bundling edit", "Does it route coordination-of-benefits and coding-edit denials to the right team?",
     "1. Open PRH-10225.\n2. Open PRH-10226.",
     "PRH-10225 (3 lines): Bill other payer → Registration (employer plan is primary). PRH-10226: Correct & resubmit → Coding "
     "(NCCI pair allows a modifier; separate session documented).", tc_other_denials),
    ("TC-11", "AI", "Explain with AI", "Does the AI explain a denial in plain English without changing the decision?",
     "1. Worklist → open the top item.\n2. Click 'Explain with AI' (the button counts seconds while it works).",
     "A plain-English explanation 'by azure:gpt-4.1-mini' appears in a few seconds. The action and queue stay the same.", tc_explain),
    ("TC-12", "Control", "Human override becomes a precedent", "Do people stay in control, and does the system learn from them?",
     "1. Open any open item.\n2. Click Override with no reason (it asks for one).\n3. Pick another action, type a reason, click Override.",
     "Without a reason it asks for one. With a reason: saved, status DONE, added to the policy library as a precedent, Resolved +1.", tc_override),
    ("TC-13", "AI", "Ask the assistant; guardrails hold", "Can a manager ask in plain English, and does the AI stay on topic?",
     "1. Assistant → 'What is due this week?'.\n2. Type 'Why was PRH-10222 denied?'.\n3. 'What is the Medicare timely filing limit?'.\n4. 'Write me a poem about cats'.",
     "Answers 1-3 with claim numbers, dollars and dates (12 months for Medicare). Politely declines the poem.", tc_assistant),
    ("TC-14", "Ingestion", "View raw file; bad file is quarantined", "Is a broken file handled safely?",
     "1. Remittance inbox → 'view raw' on a live file.\n2. Upload any non-X12 text file in 'Upload a file'.",
     "Raw view shows the X12 text (starts with ISA). The bad file is QUARANTINED with a reason; nothing breaks.", tc_files),
    ("TC-15", "Quality", "Proof & quality", "Can we prove the decisions are right and safe?",
     "1. Open Proof & quality.",
     "Denial decisions correct ≥95% (expected 100%), noise filtered ≥99%, 0 unsafe actions, ranking AUC shown.", tc_proof),
    ("TC-16", "Media", "Explainer page, GIF and MP4", "Can people watch the story without a live demo?",
     "1. Click 'How it works ▸' in the sidebar.\n2. Play; use Download MP4 / Download GIF.",
     "Explainer plays in the browser; GIF and MP4 download and play.", tc_explainer),
    ("TC-17", "Setup", "Leave the app clean", "Is the app ready for the next demo?",
     "1. Remittance inbox → Reset demo data → OK.",
     "Resolved 0, no Teams alerts, live files ready to ingest again.", tc_clean),
]


def write_workbook(path: str, rows: list[dict], meta: dict) -> None:
    wb = Workbook()
    thin = Side(style="thin", color="E0E3E7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    base = Font(name=FONT, size=10)
    head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    yellow = PatternFill("solid", fgColor="FFF6DB")

    ws = wb.active
    ws.title = "Test Cases"
    ws["A1"] = "Denial Navigator · Business test cases"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    ws["A2"] = (f"Environment: {meta['url']}  ·  Version: {meta['build']}  ·  AI: {meta['mode']}  ·  Automated run: {meta['run_at']} UTC  ·  "
                "All data synthetic (no patient data)")
    ws["A2"].font = Font(name=FONT, size=9, color="52606D")
    ws["A3"] = ("Columns A-I were filled in by the automated run. To validate: follow 'How to test' in the web app, compare with "
                "'Expected result', then fill the yellow columns (choose PASS or FAIL).")
    ws["A3"].font = Font(name=FONT, size=9, italic=True, color="52606D")
    headers = ["ID", "Area", "Test", "Business question", "How to test (in the web app)", "Expected result",
               "Automated result (captured)", "Automated status", "Response time (s)",
               "Your result (PASS/FAIL)", "Tester", "Date", "Notes"]
    widths = [8, 11, 26, 30, 44, 50, 70, 11, 10, 13, 14, 12, 28]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=5, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = head, PatternFill("solid", fgColor=NAVY if i < 10 else BLUE), Alignment(wrap_text=True, vertical="center"), border
        ws.column_dimensions[c.column_letter].width = w
    for r, row in enumerate(rows, 6):
        vals = [row["id"], row["area"], row["title"], row["question"], row["steps"], row["expected"], row["actual"],
                row["status"], round(row["seconds"], 1), "", "", "", ""]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.alignment, c.border = base, wrap, border
            if i >= 10:
                c.fill = yellow
        ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True, color=BLUE)
        for col in (8, 10):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center", vertical="top")
        ws.row_dimensions[r].height = 140
    last = 5 + len(rows)
    dv = DataValidation(type="list", formula1='"PASS,FAIL"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J6:J{last}")
    for rng in (f"H6:H{last}", f"J6:J{last}"):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=PatternFill("solid", fgColor="E3F3EA"),
                                                      font=Font(name=FONT, bold=True, color="1E7A50")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=PatternFill("solid", fgColor="FCE4DD"),
                                                      font=Font(name=FONT, bold=True, color=CORAL)))
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"A5:M{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:5"

    sm = wb.create_sheet("Summary")
    sm["A1"] = "Test run summary"
    sm["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    lines = [("Test cases", f"=COUNTA('Test Cases'!A6:A{last})"),
             ("Automated: passed", f"=COUNTIF('Test Cases'!H6:H{last},\"PASS\")"),
             ("Automated: failed", f"=COUNTIF('Test Cases'!H6:H{last},\"FAIL\")"),
             ("Automated pass rate", "=IF(B3=0,0,B4/B3)"),
             ("Validated by a person: passed", f"=COUNTIF('Test Cases'!J6:J{last},\"PASS\")"),
             ("Validated by a person: failed", f"=COUNTIF('Test Cases'!J6:J{last},\"FAIL\")"),
             ("Not yet validated", f"=B3-B7-B8"),
             ("Average response time (s)", f"=AVERAGE('Test Cases'!I6:I{last})"),
             ("AI mode", meta["mode"]), ("Version", meta["build"]), ("Environment", meta["url"]), ("Automated run at (UTC)", meta["run_at"])]
    for i, (k, v) in enumerate(lines, 3):
        sm.cell(row=i, column=1, value=k).font = Font(name=FONT, size=10, bold=True)
        sm.cell(row=i, column=2, value=v).font = base
    sm["B6"].number_format = "0%"
    sm["B10"].number_format = "0.0"
    sm.column_dimensions["A"].width = 30
    sm.column_dimensions["B"].width = 70

    how = wb.create_sheet("How to validate")
    steps = [
        "How to validate Denial Navigator (about 20 minutes)",
        "1. Open the app link in Chrome or Edge. Press Cmd+Shift+R (Mac) or Ctrl+Shift+R (Windows) once.",
        "2. Check the sidebar shows the same Version as this workbook (Summary sheet).",
        "3. Go through TC-01 to TC-17 in order on the 'Test Cases' sheet. Follow 'How to test', compare with 'Expected result'.",
        "4. In the yellow columns choose PASS or FAIL, add your name, the date and any notes.",
        "5. Finish with TC-17 (Reset demo data) so the app is clean for the next demo.",
        "",
        "To re-run the automated checks yourself, in Azure Cloud Shell:",
        "   cd ~/claims-denial-navigator/poc && python3 scripts/run_test_cases.py <app link>",
        "   The workbook is written to ~/claims-denial-navigator/poc/Denial_Navigator_Test_Results.xlsx",
        "   Download it: Cloud Shell toolbar → Manage files → Download → paste that path.",
        "",
        "Colours: PASS (green) = behaved as expected. FAIL (red) = did not; see the result column.",
        "All data is synthetic: a fictional hospital, fictional payers and patients. No real patient information is used.",
    ]
    for i, t in enumerate(steps, 1):
        c = how.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, size=12 if i == 1 else 10, bold=i == 1, color=NAVY if i == 1 else "1F2933")
    how.column_dimensions["A"].width = 120
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("url")
    ap.add_argument("--out", default="Denial_Navigator_Test_Results.xlsx")
    ap.add_argument("--mode-label", default=None, help="override the AI label written in the workbook")
    ap.add_argument("--env-label", default=None, help="override the environment label written in the workbook")
    a = ap.parse_args()
    app = App(a.url)
    meta = app.wait()
    ctx: dict = {}
    rows = []
    for cid, area, title, question, steps, expected, fn in CASES:
        try:
            ok, actual, secs = fn(app, ctx)
        except (KeyError, StopIteration) as exc:  # needs the remittance ingested by an earlier case
            ok, actual, secs = False, f"Not run: depends on an earlier case that did not complete ({type(exc).__name__}: {exc})", 0.0
        except Exception as exc:  # a failed step is a FAIL, not a crash
            ok, actual, secs = False, f"Error: {type(exc).__name__}: {exc}", 0.0
        rows.append({"id": cid, "area": area, "title": title, "question": question, "steps": steps, "expected": expected,
                     "actual": actual, "status": "PASS" if ok else "FAIL", "seconds": secs})
        print(f"{'✅' if ok else '❌'} {cid} {title} ({secs:.1f}s)\n   {actual}\n")
    run_meta = {"url": a.env_label or a.url, "build": ctx.get("build") or meta.get("build", "-"),
                "mode": a.mode_label or ("Azure OpenAI" if meta.get("agent_mode") == "azure" else "offline templates"),
                "run_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")}
    if not any(r["id"] == "TC-17" and r["status"] == "PASS" for r in rows):
        app.call("POST", "/api/demo/reset", timeout=300)  # always leave the app clean for a live demo
    out = os.path.abspath(a.out)
    write_workbook(out, rows, run_meta)
    passed = sum(r["status"] == "PASS" for r in rows)
    print(f"{passed} of {len(rows)} passed")
    print(f"Workbook saved: {out}  ({os.path.getsize(out):,} bytes)")
    return 0 if passed == len(rows) else 1


if __name__ == "__main__":
    raise SystemExit(main())
